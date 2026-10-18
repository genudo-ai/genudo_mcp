#!/usr/bin/env python3
"""
Formatted Excel Knowledge Base and Clean CSV generator for Shopify catalogs.
Optimized specifically for GenuDo AI agent grounding, Arabic title preservation/auto-translation,
and human review highlighting.
"""

import json
import re
import os
import csv
import time
import sys
import urllib.request
import urllib.parse
from concurrent.futures import ThreadPoolExecutor

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# CLI argument parsing
# Usage: python3 generate_active_excel.py <json_path> <xlsx_path> <csv_path> [shop_domain] [active_only] [cache_path]
JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else "output/products_clean.json"
ACTIVE_XLSX_PATH = sys.argv[2] if len(sys.argv) > 2 else "output/products_knowledge_base_active.xlsx"
ACTIVE_CSV_PATH = sys.argv[3] if len(sys.argv) > 3 else "output/products_catalog_clean.csv"
SHOP_DOMAIN = sys.argv[4] if len(sys.argv) > 4 else ""
ACTIVE_ONLY = sys.argv[5].lower() != "false" if len(sys.argv) > 5 else True
CACHE_PATH = sys.argv[6] if len(sys.argv) > 6 else os.path.join(os.path.dirname(ACTIVE_XLSX_PATH) or ".", ".cache", "arabic_titles_cache.json")

ARABIC_REGEX = re.compile(r'[\u0600-\u06FF]')


def load_translation_cache():
    if os.path.exists(CACHE_PATH):
        try:
            with open(CACHE_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_translation_cache(cache):
    try:
        os.makedirs(os.path.dirname(CACHE_PATH), exist_ok=True)
        with open(CACHE_PATH, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Notice: Could not write cache file {CACHE_PATH}: {e}")


def translate_to_arabic(text):
    if not text or not text.strip():
        return ""
    try:
        url = 'https://translate.googleapis.com/translate_a/single?' + urllib.parse.urlencode({
            'client': 'gtx',
            'sl': 'en',
            'tl': 'ar',
            'dt': 't',
            'q': text.strip()
        })
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)'})
        with urllib.request.urlopen(req, timeout=10) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            translated = ''.join(part[0] for part in data[0] if part and part[0])
            return translated.strip()
    except Exception as e:
        print(f"Notice: Translation fallback for '{text[:25]}...': {e}")
        return text.strip()


def get_arabic_title(product, cache):
    title = (product.get("title") or "").strip()
    pid = str(product.get("numeric_product_id") or product.get("shopify_product_id", "").split("/")[-1])

    if ARABIC_REGEX.search(title):
        return title, False

    if pid in cache:
        return cache[pid], True
    if title in cache:
        return cache[title], True

    translated = translate_to_arabic(title)
    cache[pid] = translated
    return translated, True


def batch_translate_missing(products, cache):
    needed = []
    for p in products:
        title = (p.get("title") or "").strip()
        pid = str(p.get("numeric_product_id") or p.get("shopify_product_id", "").split("/")[-1])
        if not ARABIC_REGEX.search(title) and pid not in cache and title not in cache:
            needed.append((pid, title))

    if not needed:
        return

    print(f"Translating {len(needed)} English product titles to Arabic for agent grounding...")

    def worker(item):
        pid, title = item
        res = translate_to_arabic(title)
        time.sleep(0.05)
        return pid, res

    with ThreadPoolExecutor(max_workers=6) as executor:
        results = executor.map(worker, needed)
        for pid, translated in results:
            cache[pid] = translated

    save_translation_cache(cache)


def extract_target_age(tags, title, description):
    age_matches = []
    for tag in tags:
        if re.search(r'\b(year|years|month|months|سنة|سنوات|أشهر|شهر)\b', tag, re.I):
            age_matches.append(tag.strip())
        elif re.search(r'\b\d+[-+]\d*\s*(y|yr|years?)\b', tag, re.I):
            age_matches.append(tag.strip())

    if age_matches:
        return "; ".join(dict.fromkeys(age_matches))

    title_match = re.search(r'\(Ages?\s*([^)]+)\)', title, re.I)
    if title_match:
        return title_match.group(1).strip()

    desc_match = re.search(r'Recommended Age:\s*([^\n•\r]+)', description, re.I)
    if desc_match:
        return desc_match.group(1).strip()

    return "All Ages"


def build_discount_info(variants):
    discounts = []
    for v in variants:
        price = v.get("price") or 0.0
        compare = v.get("compare_at_price")
        if compare and compare > price:
            diff = compare - price
            pct = round((diff / compare) * 100)
            discounts.append(f"Was {compare:.2f} SAR (Save {diff:.2f} SAR - {pct}%)")

    if discounts:
        return "; ".join(dict.fromkeys(discounts))
    return "Regular Price"


def build_price_range(variants):
    prices = [v.get("price") for v in variants if v.get("price") is not None]
    if not prices:
        return ""
    min_p, max_p = min(prices), max(prices)
    if min_p == max_p:
        return f"{min_p:.2f} SAR"
    return f"{min_p:.2f} - {max_p:.2f} SAR"


def build_stock_status(variants):
    total_qty = 0
    has_trackable_qty = False
    any_available = False

    for v in variants:
        if v.get("available_for_sale"):
            any_available = True
        qty = v.get("inventory_quantity")
        if qty is not None:
            has_trackable_qty = True
            total_qty += max(0, qty)

    if not any_available or (has_trackable_qty and total_qty <= 0):
        return "OUT_OF_STOCK (0 units)"

    if has_trackable_qty:
        return f"IN_STOCK ({total_qty} units)"
    return "IN_STOCK (Available)"


def build_variants_summary(variants):
    if not variants:
        return "No Options (Standard Item)"

    # Case 1: Single Standalone Product
    if len(variants) == 1:
        v = variants[0]
        title = (v.get("variant_title") or "").strip()
        sku = v.get("sku") or "NO-SKU"

        if title.lower() in ["default title", "default", ""]:
            return f"No Options (Standard Item) | SKU: {sku}"
        else:
            return f"Option [{title}] | SKU: {sku}"

    # Case 2: Multi-Variant Product
    option_names = set()
    for v in variants:
        for opt_k in (v.get("selected_options") or {}).keys():
            if opt_k.lower() != "title":
                option_names.add(opt_k)

    prefix = f"{', '.join(option_names)} Options:\n" if option_names else "Multiple Options:\n"

    parts = []
    for v in variants:
        title = v.get("variant_title") or "Option"
        sku = v.get("sku") or "NO-SKU"
        qty = v.get("inventory_quantity")
        avail = v.get("available_for_sale")

        opts = v.get("selected_options") or {}
        opt_desc = ", ".join(f"{val}" for k, val in opts.items() if k.lower() != "title")
        label = opt_desc if opt_desc else title

        stock_str = f"In Stock: {qty} units" if qty is not None and qty > 0 else ("In Stock" if avail else "Sold Out")
        parts.append(f"• {label} | SKU: {sku} ({stock_str})")

    return prefix + "\n".join(parts)


def run_export():
    if not os.path.exists(JSON_PATH):
        raise FileNotFoundError(f"Missing {JSON_PATH}. Run product extraction first.")

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        all_products = json.load(f)

    if ACTIVE_ONLY:
        target_products = [p for p in all_products if p.get("status") == "ACTIVE"]
    else:
        target_products = all_products

    print(f"Total catalog products: {len(all_products)}")
    print(f"Selected products for knowledge base: {len(target_products)} (active_only={ACTIVE_ONLY})")

    cache = load_translation_cache()
    batch_translate_missing(target_products, cache)

    headers = [
        ("default_id", "Stable ID", 22),
        ("product_name", "Product Name (Original)", 42),
        ("arabic_title", "Arabic Title (اسم المنتج بالعربي)", 42),
        ("target_age", "Target Age", 18),
        ("category_and_tags", "Category & Tags", 34),
        ("stock_status", "Stock Status", 20),
        ("price_range", "Price Range", 16),
        ("discount_info", "Discount / Promotion", 24),
        ("description", "Clean Description & Sales Copy", 65),
        ("variants_summary", "Options & SKU (خيارات المنتج)", 45),
        ("image_url", "Featured Image URL", 45),
        ("product_url", "Public Store Link", 45),
        ("admin_url", "Shopify Admin Link", 45),
    ]

    csv_rows = []
    csv_header = [h[0] for h in headers]
    csv_rows.append(csv_header)

    clean_domain = SHOP_DOMAIN.replace("https://", "").replace("http://", "").strip().rstrip("/")

    data_rows = []
    auto_translated_flags = []

    for p in target_products:
        pid = str(p.get("numeric_product_id") or p.get("shopify_product_id", "").split("/")[-1])
        default_id = f"prod-{pid}"
        title = p.get("title") or ""
        desc = p.get("description") or ""
        category = p.get("category") or p.get("product_type") or "General"
        tags = p.get("tags") or []
        variants = p.get("variants") or []

        arabic_title, is_auto = get_arabic_title(p, cache)
        auto_translated_flags.append(is_auto)

        target_age = extract_target_age(tags, title, desc)
        tags_str = ", ".join(tags[:6])
        category_and_tags = f"Category: {category} | Tags: {tags_str}" if tags_str else f"Category: {category}"
        stock_status = build_stock_status(variants)
        price_range = build_price_range(variants)
        discount_info = build_discount_info(variants)
        variants_summary = build_variants_summary(variants)
        image_url = p.get("featured_image_url") or ""
        product_url = p.get("product_url") or ""

        admin_url = p.get("admin_url") or (f"https://{clean_domain}/admin/products/{pid}" if clean_domain and pid else "")

        row_data = [
            default_id,
            title,
            arabic_title,
            target_age,
            category_and_tags,
            stock_status,
            price_range,
            discount_info,
            desc,
            variants_summary,
            image_url,
            product_url,
            admin_url,
        ]

        data_rows.append(row_data)
        csv_rows.append(row_data)

    # 1. Save Clean CSV with UTF-8 BOM
    os.makedirs(os.path.dirname(os.path.abspath(ACTIVE_CSV_PATH)), exist_ok=True)
    with open(ACTIVE_CSV_PATH, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for r in csv_rows:
            writer.writerow(r)
    print(f"Successfully saved clean CSV with UTF-8 BOM: {ACTIVE_CSV_PATH}")

    # 2. Save Formatted Excel Knowledge Base (if openpyxl is available)
    if not OPENPYXL_AVAILABLE:
        print("Notice: openpyxl is not installed. Skipped Excel styling; CSV generated successfully.")
        return

    os.makedirs(os.path.dirname(os.path.abspath(ACTIVE_XLSX_PATH)), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Active Products Catalog" if ACTIVE_ONLY else "Products Catalog"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    fill_review_yellow = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    border_review_yellow = Border(
        left=Side(style="thin", color="F59E0B"),
        right=Side(style="thin", color="F59E0B"),
        top=Side(style="thin", color="F59E0B"),
        bottom=Side(style="thin", color="F59E0B"),
    )

    for col_idx, (col_id, col_label, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border_thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 36

    zebra_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    zebra_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_link = Font(name="Calibri", size=10, color="1D4ED8", underline="single")
    font_in_stock = Font(name="Calibri", size=10, bold=True, color="166534")
    font_out_stock = Font(name="Calibri", size=10, bold=True, color="991B1B")
    fill_in_stock = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_out_stock = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    for i, row_data in enumerate(data_rows):
        row_idx = i + 2
        is_auto = auto_translated_flags[i]
        row_fill = zebra_even if row_idx % 2 == 0 else zebra_odd

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            cell.fill = row_fill
            cell.font = font_regular

            is_empty = val is None or str(val).strip() == ""

            # Horizontal alignment
            if col_idx in [1, 4, 7]:  # default_id, target_age, price_range
                h_align = "center"
            elif col_idx == 6:        # stock_status
                h_align = "center"
            elif col_idx == 3:        # arabic_title (Right-to-Left alignment)
                h_align = "right"
            else:
                h_align = "left"

            cell.alignment = Alignment(horizontal=h_align, vertical="top", wrap_text=True)

            if col_idx == 1:
                cell.font = font_bold
            elif col_idx == 6:  # stock_status
                if "IN_STOCK" in str(val):
                    cell.fill = fill_in_stock
                    cell.font = font_in_stock
                else:
                    cell.fill = fill_out_stock
                    cell.font = font_out_stock
            elif col_idx in [12, 13] and val and str(val).startswith("http"):  # product_url, admin_url
                cell.font = font_link
                cell.hyperlink = str(val)

            # Soft yellow highlights for auto-translated Arabic titles or empty values requiring human review
            if col_idx == 3 and is_auto:
                cell.fill = fill_review_yellow
                cell.border = border_review_yellow
            elif is_empty:
                cell.fill = fill_review_yellow
                cell.border = border_review_yellow

        ws.row_dimensions[row_idx].height = 55

    ws.freeze_panes = "A2"
    wb.save(ACTIVE_XLSX_PATH)
    print(f"Successfully saved formatted Excel: {ACTIVE_XLSX_PATH}")


if __name__ == "__main__":
    run_export()
