#!/usr/bin/env python3
"""Allocate agent-QA stories into 2-week sprints and write an Excel backlog.

Usage:
  python3 build_backlog.py stories.json -o sprint_backlog.xlsx [--capacity 35]
  python3 build_backlog.py --self-check

stories.json is a list of objects:
  {"id": "BUG-01", "priority": "P0", "category": "Order journey",
   "story": "As a buyer ...", "component": "Stage instruction", "points": 5,
   "acceptance": "...", "linked_test": "J01", "status": "To Do"}
"""
import argparse
import json
import sys

FIBONACCI = {1, 2, 3, 5, 8, 13}
PRIORITY_ORDER = {"P0": 0, "P1": 1, "P2": 2}
COLUMNS = ["id", "priority", "category", "story", "component", "points",
           "acceptance", "linked_test", "status"]


def allocate(stories, capacity):
    """Return a list of sprints (lists of stories), P0 first, then P1, P2.

    P0 always lands in sprint 1 even past capacity, so blockers are never deferred.
    """
    for s in stories:
        if s.get("priority") not in PRIORITY_ORDER:
            raise ValueError(f"{s.get('id')}: priority must be P0, P1 or P2")
        if s.get("points") not in FIBONACCI:
            raise ValueError(f"{s.get('id')}: points must be one of {sorted(FIBONACCI)}")
    ordered = sorted(stories, key=lambda s: PRIORITY_ORDER[s["priority"]])
    sprints = [[]]
    load = [0]
    for s in ordered:
        if s["priority"] == "P0":
            sprints[0].append(s)
            load[0] += s["points"]
            continue
        # ponytail: first-fit, no bin-packing; fine for a backlog of a few dozen stories
        for i in range(len(sprints)):
            if load[i] + s["points"] <= capacity:
                break
        else:
            sprints.append([])
            load.append(0)
            i = len(sprints) - 1
        sprints[i].append(s)
        load[i] += s["points"]
    return sprints


def write_xlsx(sprints, capacity, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Sprint", "Stories", "Points", "Capacity", "P0", "Over capacity"])
    header_fill = PatternFill("solid", start_color="1B365D")
    for n, sprint in enumerate(sprints, 1):
        pts = sum(s["points"] for s in sprint)
        summary.append([f"Sprint {n}", len(sprint), pts, capacity,
                        sum(s["priority"] == "P0" for s in sprint), "YES" if pts > capacity else ""])
        ws = wb.create_sheet(f"Sprint {n}")
        ws.append([c.replace("_", " ").title() for c in COLUMNS])
        for s in sprint:
            ws.append([s.get(c, "") for c in COLUMNS])
            for cell in ws[ws.max_row]:
                # Story text can quote customer messages; never let it become a live formula.
                if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
                    cell.data_type = "s"
        dv = DataValidation(type="list", formula1='"To Do,In Progress,Verified,Done"')
        ws.add_data_validation(dv)
        dv.add(f"I2:I{max(len(sprint) + 1, 2)}")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for col, width in zip("ABCDEFGHI", (10, 8, 18, 60, 20, 8, 50, 12, 12)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    wb.save(path)


def self_check():
    def st(i, p, pts):
        return {"id": i, "priority": p, "points": pts}
    sprints = allocate([st("a", "P2", 8), st("b", "P0", 13), st("c", "P1", 13),
                        st("d", "P0", 13), st("e", "P1", 5)], capacity=30)
    assert [s["id"] for s in sprints[0]] == ["b", "d"], sprints  # P0 first, 26 pts
    assert [s["id"] for s in sprints[1]] == ["c", "e", "a"], sprints  # 13 won't fit in sprint 1
    # P0 overflow stays in sprint 1
    assert len(allocate([st("x", "P0", 13)] * 4, capacity=30)) == 1
    for bad in (st("y", "P3", 5), st("z", "P0", 4)):
        try:
            allocate([bad], 30)
            raise AssertionError("expected ValueError")
        except ValueError:
            pass
    import os
    import tempfile
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "b.xlsx")
        evil = {"id": "e", "priority": "P0", "points": 1, "story": '=HYPERLINK("http://x","y")'}
        write_xlsx(allocate([evil], 30), 30, path)
        cell = load_workbook(path)["Sprint 1"]["D2"]
        assert cell.data_type == "s" and cell.value.startswith("="), (cell.data_type, cell.value)
    print("self-check ok")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("stories", nargs="?")
    ap.add_argument("-o", "--output", default="sprint_backlog.xlsx")
    ap.add_argument("--capacity", type=int, default=35)
    ap.add_argument("--self-check", action="store_true")
    a = ap.parse_args()
    if a.self_check:
        self_check()
        sys.exit(0)
    if not a.stories:
        ap.error("stories.json is required")
    with open(a.stories, encoding="utf-8") as f:
        sprints = allocate(json.load(f), a.capacity)
    write_xlsx(sprints, a.capacity, a.output)
    print(f"wrote {a.output}: {len(sprints)} sprint(s)")
